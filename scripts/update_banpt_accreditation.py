import csv
import json
import os
import re
import time
from collections import Counter
from datetime import date, datetime
from pathlib import Path

import requests
import urllib3
from openpyxl import load_workbook

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

ROOT = Path(__file__).resolve().parents[1]
CSV_PATH = ROOT / "database" / "seeders" / "data" / "data_akreditasi_lengkap.csv"
CSV_SOURCE_PATH = Path(os.environ.get("AKREDITASI_SOURCE_CSV", CSV_PATH))
REPORT_PATH = ROOT / "database" / "seeders" / "data" / "data_akreditasi_lengkap_update_report.md"
CACHE_PATH = ROOT / ".tmp" / "banpt_accreditation_cache.json"
SCOPE_XLSX = Path(r"l:\Husni\Project\Daisy\Dokumen\Prodi\260526 - Cakupan Prodi LAMDEPILAR.xlsx")
BASE_URL = "https://service.banpt.or.id/bianglala/"
TODAY = date(2026, 5, 30)


def norm(value):
    return re.sub(r"\s+", " ", str(value or "").strip()).casefold()


def comparable_program_name(value):
    name = norm(value)
    name = re.sub(r"\s*\([^)]*\)\s*", " ", name)
    name = re.sub(r"^(magister|doktor)\s+", "", name)
    name = name.replace("&", " dan ")
    name = re.sub(r"\bsumber\s*daya\b", "sumberdaya", name)
    name = re.sub(r"\blansekap\b", "lanskap", name)
    name = re.sub(r"\bperencanaan pariwisata\b", "perencanaan dan pengembangan pariwisata", name)
    name = re.sub(r"\s+", " ", name).strip()
    return name


def has_campus_suffix(value):
    return bool(re.search(r"\([^)]*kampus[^)]*\)", str(value or ""), re.IGNORECASE))


def has_degree_prefix(value):
    return bool(re.match(r"^\s*(magister|doktor)\s+", str(value or ""), re.IGNORECASE))


def strip_degree_prefix(value):
    return re.sub(r"^\s*(magister|doktor)\s+", "", str(value or "").strip(), flags=re.IGNORECASE).strip()


def in_scope(program_name, scope_names):
    comparable = comparable_program_name(program_name)
    if comparable in scope_names:
        return True
    # Extra cautious fallback for official BAN-PT names that contain a valid scoped
    # program plus campus/location suffixes.
    return any(comparable.startswith(scope + " ") for scope in scope_names)


def parse_ps(value):
    match = re.match(r"^\s*(\d+)\s+-\s+(.+?)\s*$", value or "")
    if not match:
        return None
    rest = match.group(2).strip()
    jenjang = None
    name = None
    for prefix in ("S3 Terapan", "S2 Terapan", "D4", "D-IV", "D3", "D-III", "Profesi", "S3", "S2", "S1"):
        if norm(rest).startswith(norm(prefix) + " "):
            jenjang = prefix
            name = rest[len(prefix) :].strip()
            break
    if not jenjang:
        parts = rest.split(maxsplit=1)
        if len(parts) != 2:
            return None
        jenjang, name = parts
    return {
        "raw": value,
        "code": match.group(1),
        "jenjang": jenjang,
        "name": name,
    }


def expiry_status(value):
    if not value:
        return "Tidak Ada Data"
    try:
        expiry = datetime.strptime(value, "%Y-%m-%d").date()
    except ValueError:
        return "Tidak Ada Data"
    remaining = (expiry - TODAY).days
    if remaining <= 0:
        return "Kadaluarsa"
    if remaining <= 180:
        return f"{remaining} hari lagi kadaluarsa"
    return "Masih Berlaku"


def load_scope():
    workbook = load_workbook(SCOPE_XLSX, read_only=True, data_only=True)
    names = []
    by_sheet = {}
    for sheet in workbook.worksheets:
        sheet_names = []
        for row in sheet.iter_rows(min_row=3, values_only=True):
            if len(row) > 2 and row[2] and str(row[2]).strip():
                name = str(row[2]).strip()
                names.append(name)
                sheet_names.append(name)
        by_sheet[sheet.title] = sheet_names
    return names, by_sheet


def load_cache():
    if not CACHE_PATH.exists():
        return {}
    return json.loads(CACHE_PATH.read_text(encoding="utf-8"))


def save_cache(cache):
    CACHE_PATH.parent.mkdir(parents=True, exist_ok=True)
    CACHE_PATH.write_text(json.dumps(cache, ensure_ascii=False, indent=2), encoding="utf-8")


class BanptClient:
    def __init__(self, cache):
        self.cache = cache
        self.session = requests.Session()
        self.session.headers.update({"User-Agent": "Mozilla/5.0"})

    def get_json(self, endpoint, params):
        key = endpoint + "?" + "&".join(f"{k}={params[k]}" for k in sorted(params))
        if key in self.cache:
            return self.cache[key]
        response = self.session.get(BASE_URL + endpoint, params=params, timeout=30, verify=False)
        response.raise_for_status()
        try:
            data = response.json()
        except ValueError:
            data = None
        self.cache[key] = data
        time.sleep(0.05)
        return data

    def search_pt_candidates(self, name):
        try:
            data = self.get_json("searchPT.php", {"term": name}) or []
        except requests.HTTPError:
            if "'" not in name:
                raise
            data = self.get_json("searchPT.php", {"term": name.replace("'", "")}) or []
        items = [item for item in data if item]
        if not items and "'" in name:
            data = self.get_json("searchPT.php", {"term": name.replace("'", "")}) or []
            items = [item for item in data if item]
        if not items:
            return []
        needle = norm(name)
        parsed = []
        for item in items:
            label = re.sub(r"^\d+\s+-\s+", "", item).split(",")[0].strip()
            parsed.append((item, label))
        ordered = []
        for item, label in parsed:
            if norm(label) == needle and item not in ordered:
                ordered.append(item)
        for item, label in parsed:
            if (needle in norm(label) or norm(label) in needle) and item not in ordered:
                ordered.append(item)
        for item in items:
            if item not in ordered:
                ordered.append(item)
        return ordered

    def search_ps(self, pt, prodi, jenjang):
        terms = [prodi, re.sub(r"\s*\([^)]*\)\s*", " ", prodi).strip(), re.sub(r"^(Magister|Doktor)\s+", "", prodi).strip()]
        if "Lansekap" in prodi:
            terms.append(prodi.replace("Lansekap", "Lanskap"))
        if "&" in prodi:
            terms.append(prodi.replace("&", "dan"))
        terms = list(dict.fromkeys(term for term in terms if term))
        candidates = []
        for term in terms:
            data = self.get_json("searchPS.php", {"pt": pt, "term": term}) or []
            candidates.extend(parsed for parsed in (parse_ps(item) for item in data if item) if parsed)
        for candidate in candidates:
            if comparable_program_name(candidate["name"]) == comparable_program_name(prodi) and norm(candidate["jenjang"]) == norm(jenjang):
                return candidate["raw"]
        for candidate in candidates:
            if comparable_program_name(candidate["name"]) == comparable_program_name(prodi):
                return candidate["raw"]
        return None

    def accreditation(self, pt, ps):
        return self.get_json("riakps.php", {"pt": pt, "ps": ps})


def active_record(records):
    if not records:
        return None
    active = [record for record in records if len(record) >= 7 and norm(record[6]) == "ya"]
    return active[-1] if active else None


def main():
    scope_names, scope_by_sheet = load_scope()
    scope_norm = {comparable_program_name(name) for name in scope_names}
    cache = load_cache()
    client = BanptClient(cache)

    with CSV_SOURCE_PATH.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        rows = list(reader)
        fieldnames = reader.fieldnames

    updated_rows = []
    removed_rows = []
    inactive_rows = []
    failures = []
    diffs = []
    summary = Counter()

    for index, row in enumerate(rows, start=2):
        if not in_scope(row["Program Studi"], scope_norm):
            removed_rows.append(row)
            summary["removed_out_of_scope"] += 1
            continue

        summary["kept_in_scope"] += 1
        new_row = dict(row)
        reason = None
        try:
            pt_candidates = client.search_pt_candidates(row["Universitas"])
            if not pt_candidates:
                reason = "Perguruan tinggi tidak ditemukan di BAN-PT"
            else:
                pt = None
                ps = None
                for candidate in pt_candidates:
                    ps = client.search_ps(candidate, row["Program Studi"], row["Jenjang"])
                    if ps:
                        pt = candidate
                        break
                if not ps:
                    reason = "Program studi/jenjang tidak ditemukan di BAN-PT"
                else:
                    record = active_record(client.accreditation(pt, ps))
                    if not record:
                        inactive_rows.append(row)
                        summary["removed_inactive"] += 1
                        continue
                    else:
                        new_values = {
                            "Peringkat_Akreditasi": record[4],
                            "Tanggal_Kadaluarsa": record[5],
                            "Status_Kadaluarsa": expiry_status(record[5]),
                        }
                        if has_degree_prefix(row["Program Studi"]) and not has_campus_suffix(row["Program Studi"]):
                            new_values["Program Studi"] = strip_degree_prefix(row["Program Studi"])
                            new_values["Jenjang"] = record[2]
                        if record[3] and re.match(r"\d{4}-", record[3]):
                            new_values["Tahun"] = record[3][:4]
                        changed = {
                            key: (new_row.get(key, ""), value)
                            for key, value in new_values.items()
                            if str(new_row.get(key, "")) != str(value)
                        }
                        new_row.update(new_values)
                        if changed:
                            diffs.append(
                                {
                                    "line": index,
                                    "universitas": row["Universitas"],
                                    "program_studi": row["Program Studi"],
                                    "jenjang": row["Jenjang"],
                                    "changes": changed,
                                }
                            )
                            summary["updated"] += 1
                        else:
                            summary["unchanged"] += 1
        except Exception as exc:
            reason = f"Error saat akses BAN-PT: {type(exc).__name__}: {exc}"

        if reason:
            failures.append(
                {
                    "line": index,
                    "universitas": row["Universitas"],
                    "program_studi": row["Program Studi"],
                    "jenjang": row["Jenjang"],
                    "reason": reason,
                }
            )
            summary["not_verified"] += 1
        updated_rows.append(new_row)

        if index % 50 == 0:
            save_cache(cache)
            print(f"Processed line {index}: {dict(summary)}", flush=True)

    save_cache(cache)

    with CSV_PATH.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(updated_rows)

    report = []
    report.append("# Laporan Update Data Akreditasi Lengkap")
    report.append("")
    report.append(f"- Tanggal update: {TODAY.isoformat()}")
    report.append(f"- Sumber cakupan: `{SCOPE_XLSX}`")
    report.append("- Sumber akreditasi: `https://service.banpt.or.id/bianglala/` (`searchPT.php`, `searchPS.php`, `riakps.php`)")
    report.append("- Catatan: endpoint BAN-PT Bianglala tidak mengembalikan Nomor SK, jadi kolom `Nomor SK` dipertahankan dari data lama.")
    empty_expiry_count = sum(1 for row in updated_rows if not row["Tanggal_Kadaluarsa"])
    if empty_expiry_count:
        report.append(f"- Catatan: {empty_expiry_count} baris tidak memiliki tanggal kedaluwarsa dari endpoint BAN-PT, sehingga `Tanggal_Kadaluarsa` dikosongkan dan `Status_Kadaluarsa` diisi `Tidak Ada Data`.")
    report.append("")
    report.append("## Ringkasan")
    report.append("")
    report.append(f"- Jumlah baris sebelum update: {len(rows)}")
    report.append(f"- Jumlah baris setelah update: {len(updated_rows)}")
    report.append(f"- Jumlah baris dihapus karena prodi tidak masuk cakupan Excel: {len(removed_rows)}")
    report.append(f"- Jumlah baris dihapus karena prodi tidak aktif di BAN-PT: {len(inactive_rows)}")
    report.append(f"- Jumlah baris diperbarui dari BAN-PT: {len(diffs)}")
    report.append(f"- Jumlah baris tidak berubah: {summary['unchanged']}")
    report.append(f"- Jumlah baris belum terverifikasi: {len(failures)}")
    report.append("")
    report.append("## Cakupan Prodi Excel")
    report.append("")
    for sheet, names in scope_by_sheet.items():
        report.append(f"- {sheet}: {len(names)} prodi")
    report.append("")
    report.append("## Baris Tidak Terverifikasi")
    report.append("")
    if failures:
        report.append("| Line | Universitas | Program Studi | Jenjang | Alasan |")
        report.append("| --- | --- | --- | --- | --- |")
        for item in failures:
            report.append(
                f"| {item['line']} | {item['universitas']} | {item['program_studi']} | {item['jenjang']} | {item['reason']} |"
            )
    else:
        report.append("Tidak ada.")
    report.append("")
    report.append("## Baris Dihapus Karena Di Luar Cakupan")
    report.append("")
    if removed_rows:
        report.append("| Universitas | Program Studi | Jenjang |")
        report.append("| --- | --- | --- |")
        for row in removed_rows:
            report.append(f"| {row['Universitas']} | {row['Program Studi']} | {row['Jenjang']} |")
    else:
        report.append("Tidak ada.")
    report.append("")
    report.append("## Baris Dihapus Karena Prodi Tidak Aktif Di BAN-PT")
    report.append("")
    if inactive_rows:
        report.append("| Universitas | Program Studi | Jenjang |")
        report.append("| --- | --- | --- |")
        for row in inactive_rows:
            report.append(f"| {row['Universitas']} | {row['Program Studi']} | {row['Jenjang']} |")
    else:
        report.append("Tidak ada.")
    report.append("")
    report.append("## Contoh Perubahan")
    report.append("")
    if diffs:
        report.append("| Line | Universitas | Program Studi | Jenjang | Perubahan |")
        report.append("| --- | --- | --- | --- | --- |")
        for item in diffs[:100]:
            changes = "; ".join(f"{key}: `{old}` -> `{new}`" for key, (old, new) in item["changes"].items())
            report.append(
                f"| {item['line']} | {item['universitas']} | {item['program_studi']} | {item['jenjang']} | {changes} |"
            )
    else:
        report.append("Tidak ada.")
    REPORT_PATH.write_text("\n".join(report) + "\n", encoding="utf-8")

    print(json.dumps(dict(summary), ensure_ascii=False, indent=2), flush=True)


if __name__ == "__main__":
    main()
