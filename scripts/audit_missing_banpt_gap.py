import csv
import json
import os
import re
import time
from collections import Counter, defaultdict
from pathlib import Path

import requests
import urllib3
from openpyxl import load_workbook

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

ROOT = Path(__file__).resolve().parents[1]
DBX = Path(r"l:\Husni\Project\Daisy\Dokumen\Prodi\260526 - Database Prodi.xlsx")
CAK = Path(r"l:\Husni\Project\Daisy\Dokumen\Prodi\260526 - Cakupan Prodi LAMDEPILAR.xlsx")
CSV_PATH = ROOT / "database" / "seeders" / "data" / "data_akreditasi_lengkap.csv"
REPORT_PATH = ROOT / "database" / "seeders" / "data" / "banpt_pddikti_gap_report.md"
CACHE_PATH = ROOT / ".tmp" / "banpt_gap_audit_cache.json"
BASE = "https://service.banpt.or.id/bianglala/"


def norm(value):
    return re.sub(r"\s+", " ", str(value or "").strip()).casefold()


def comp(value):
    value = norm(value)
    value = re.sub(r"\s*\([^)]*\)\s*", " ", value)
    value = re.sub(r"^(magister|doktor)\s+", "", value)
    value = value.replace("&", " dan ")
    value = re.sub(r"\bsumber\s*daya\b", "sumberdaya", value)
    value = re.sub(r"\blansekap\b", "lanskap", value)
    value = re.sub(r"\bperencanaan pariwisata\b", "perencanaan dan pengembangan pariwisata", value)
    return re.sub(r"\s+", " ", value).strip()


def as_int(value):
    if isinstance(value, (int, float)):
        return int(value)
    text = str(value or "").strip()
    if not text:
        return 0
    try:
        return int(float(text))
    except ValueError:
        return 0


def jenjang_norm(value):
    compact = norm(value).replace(" ", "")
    return {
        "d-iii": "D3",
        "diii": "D3",
        "d3": "D3",
        "d-iv": "D4",
        "div": "D4",
        "d4": "D4",
        "s1": "S1",
        "s2": "S2",
        "s3": "S3",
        "s2terapan": "S2 Terapan",
        "s3terapan": "S3 Terapan",
        "profesi": "Profesi",
    }.get(compact, str(value).strip())


def parse_ps(value):
    match = re.match(r"^\s*(\d+)\s+-\s+(.+?)\s*$", value or "")
    if not match:
        return None
    rest = match.group(2).strip()
    for prefix in ("S3 Terapan", "S2 Terapan", "D4", "D-IV", "D3", "D-III", "Profesi", "S3", "S2", "S1"):
        if norm(rest).startswith(norm(prefix) + " "):
            return {
                "raw": value,
                "code": match.group(1),
                "jenjang": jenjang_norm(prefix),
                "name": rest[len(prefix) :].strip(),
            }
    parts = rest.split(maxsplit=1)
    if len(parts) != 2:
        return None
    return {"raw": value, "code": match.group(1), "jenjang": jenjang_norm(parts[0]), "name": parts[1].strip()}


def active_record(records):
    if not records:
        return None
    active = [record for record in records if len(record) >= 7 and norm(record[6]) == "ya"]
    return active[-1] if active else None


def load_json(path):
    if not path.exists():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def save_json(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


class Banpt:
    def __init__(self, cache):
        self.cache = cache
        self.session = requests.Session()
        self.session.headers.update({"User-Agent": "Mozilla/5.0"})

    def get(self, endpoint, params):
        key = endpoint + "?" + "&".join(f"{k}={params[k]}" for k in sorted(params))
        if key in self.cache:
            return self.cache[key]
        response = self.session.get(BASE + endpoint, params=params, timeout=30, verify=False)
        response.raise_for_status()
        try:
            data = response.json()
        except ValueError:
            data = None
        self.cache[key] = data
        time.sleep(0.03)
        return data

    def pt_candidates(self, name):
        terms = [name]
        if "'" in name:
            terms.append(name.replace("'", ""))
        items = []
        for term in terms:
            try:
                data = self.get("searchPT.php", {"term": term}) or []
            except requests.HTTPError:
                continue
            for item in data:
                if item and item not in items:
                    items.append(item)
        needle = norm(name)
        parsed = []
        for item in items:
            label = re.sub(r"^\d+\s+-\s+", "", item).split(",")[0].strip()
            score = 0
            if norm(label) == needle:
                score = 2
            elif needle in norm(label) or norm(label) in needle:
                score = 1
            parsed.append((score, item))
        return [item for _, item in sorted(parsed, key=lambda pair: -pair[0])]

    def search_ps(self, pt, prodi):
        terms = [prodi]
        if comp(prodi) != norm(prodi):
            terms.append(comp(prodi))
        if "Profesi Arsitek" in prodi:
            terms.append("Pendidikan Profesi Arsitek")
        if "Pendidikan Profesi Arsitek" in prodi:
            terms.append("Profesi Arsitek")
        terms = list(dict.fromkeys(term for term in terms if term))
        candidates = []
        for term in terms:
            data = self.get("searchPS.php", {"pt": pt, "term": term}) or []
            for item in data:
                parsed = parse_ps(item)
                if parsed and parsed not in candidates:
                    candidates.append(parsed)
        return candidates

    def accreditation(self, pt, ps):
        return self.get("riakps.php", {"pt": pt, "ps": ps})


def load_scope():
    workbook = load_workbook(CAK, read_only=True, data_only=True)
    scope = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(min_row=3, values_only=True):
            if len(row) > 2 and row[2] and str(row[2]).strip():
                scope.append((sheet.title, str(row[2]).strip(), comp(row[2])))
    return scope


def find_scope(scope, prodi):
    key = comp(prodi)
    for sheet, name, scope_key in scope:
        if key == scope_key:
            return sheet, name, scope_key
    matches = [(sheet, name, scope_key) for sheet, name, scope_key in scope if key.startswith(scope_key + " ")]
    return max(matches, key=lambda row: len(row[2])) if matches else None


def expected_actual(scope):
    cols = [("S1", 4), ("S2", 5), ("S3", 6), ("D4", 7), ("D3", 8), ("D2", 9), ("D1", 10), ("Profesi", 11)]
    expected = Counter()
    extra = []
    workbook = load_workbook(DBX, read_only=True, data_only=True)
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(min_row=3, values_only=True):
            nomor = row[0] if len(row) > 0 else None
            name = row[2] if len(row) > 2 else None
            if not (isinstance(nomor, (int, float)) and isinstance(name, str) and name.strip()):
                continue
            found = find_scope(scope, name)
            values = {level: as_int(row[idx] if len(row) > idx else None) for level, idx in cols}
            if not found:
                extra.append((sheet.title, name.strip(), sum(values.values()), values))
                continue
            for level, count in values.items():
                if count:
                    expected[(found[0], found[2], level)] += count

    actual = Counter()
    existing_keys = set()
    universities_by_sheet = defaultdict(set)
    with CSV_PATH.open(encoding="utf-8", newline="") as handle:
        for row in csv.DictReader(handle):
            found = find_scope(scope, row["Program Studi"])
            if not found:
                continue
            level = jenjang_norm(row["Jenjang"])
            actual[(found[0], found[2], level)] += 1
            existing_keys.add((norm(row["Universitas"]), found[2], level))
            universities_by_sheet[found[0]].add(row["Universitas"])
    return expected, actual, universities_by_sheet, existing_keys, extra


def main():
    scope = load_scope()
    scope_by_key = {key: name for _, name, key in scope}
    expected, actual, universities_by_sheet, existing_keys, extra = expected_actual(scope)
    deficits = []
    for key in sorted(set(expected) | set(actual)):
        diff = expected.get(key, 0) - actual.get(key, 0)
        if diff > 0:
            deficits.append((key[0], key[1], key[2], diff))

    cache = load_json(CACHE_PATH)
    client = Banpt(cache)
    found = []
    searched = 0

    for sheet, prodi_key, level, deficit in deficits:
        prodi_name = scope_by_key[prodi_key]
        universities = sorted(universities_by_sheet[sheet])
        for university in universities:
            searched += 1
            try:
                pt_candidates = client.pt_candidates(university)
                match = None
                for pt in pt_candidates[:3]:
                    ps_candidates = client.search_ps(pt, prodi_name)
                    for ps in ps_candidates:
                        if ps["jenjang"] != level:
                            continue
                        if comp(ps["name"]) != prodi_key and not comp(ps["name"]).startswith(prodi_key + " "):
                            continue
                        record = active_record(client.accreditation(pt, ps["raw"]))
                        if record:
                            if (norm(record[0]), prodi_key, level) in existing_keys or (norm(university), prodi_key, level) in existing_keys:
                                continue
                            match = {
                                "sheet": sheet,
                                "university": university,
                                "target_prodi": prodi_name,
                                "target_jenjang": level,
                                "banpt_pt": pt,
                                "banpt_ps": ps["raw"],
                                "record": record,
                            }
                            break
                    if match:
                        break
                if match:
                    found.append(match)
            except Exception:
                pass
            if searched % 200 == 0:
                save_json(CACHE_PATH, cache)
                print(f"searched={searched} found={len(found)}", flush=True)

    save_json(CACHE_PATH, cache)
    grouped = defaultdict(list)
    for item in found:
        grouped[(item["sheet"], item["target_prodi"], item["target_jenjang"])].append(item)

    lines = [
        "# Audit Gap BAN-PT/PDDikti",
        "",
        "- PDDikti public API yang tersedia saat audit berada dalam mode limited/timeout, sehingga verifikasi massal PDDikti belum bisa diselesaikan otomatis.",
        "- Audit ini memakai daftar selisih dari `260526 - Database Prodi.xlsx` vs `data_akreditasi_lengkap.csv`, lalu mencari ulang kecocokan di BAN-PT Bianglala.",
        "- Prodi/riwayat BAN-PT dengan `Aktif = Tidak` tidak dihitung sebagai kandidat dan tidak dimasukkan ke hasil.",
        "- Jika sebuah gap ditemukan di BAN-PT, besar kemungkinan datanya juga perlu dicek ulang di PDDikti saat endpoint PDDikti kembali stabil.",
        "",
        "## Ringkasan Gap",
        "",
        f"- Kombinasi prodi/jenjang yang kurang: {len(deficits)}",
        f"- Total kekurangan berdasarkan detail cakupan: {sum(item[3] for item in deficits)}",
        f"- Temuan BAN-PT kandidat: {len(found)}",
        "",
        "## Extra Di Database Prodi Namun Bukan Cakupan",
        "",
    ]
    if extra:
        lines += ["| Sheet | Prodi | Jumlah |", "| --- | --- | --- |"]
        for sheet, name, total, _ in extra:
            lines.append(f"| {sheet} | {name} | {total} |")
    else:
        lines.append("Tidak ada.")
    lines += ["", "## Temuan BAN-PT Untuk Gap", ""]
    if grouped:
        for (sheet, prodi, level), items in sorted(grouped.items()):
            lines.append(f"### {sheet} - {prodi} - {level}")
            lines.append("")
            lines.append("| Universitas | Peringkat | Kedaluwarsa | Aktif | BAN-PT PS |")
            lines.append("| --- | --- | --- | --- | --- |")
            for item in items:
                record = item["record"]
                lines.append(
                    f"| {record[0]} | {record[4]} | {record[5]} | {record[6]} | {item['banpt_ps']} |"
                )
            lines.append("")
    else:
        lines.append("Belum ada temuan BAN-PT.")
    lines += ["## Defisit Detail", "", "| Sheet | Prodi | Jenjang | Kurang | Temuan BAN-PT |", "| --- | --- | --- | ---: | ---: |"]
    for sheet, key, level, deficit in deficits:
        prodi = scope_by_key[key]
        count = len(grouped.get((sheet, prodi, level), []))
        lines.append(f"| {sheet} | {prodi} | {level} | {deficit} | {count} |")
    REPORT_PATH.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(json.dumps({"deficits": len(deficits), "deficit_total": sum(item[3] for item in deficits), "banpt_found": len(found)}, indent=2), flush=True)


if __name__ == "__main__":
    main()
